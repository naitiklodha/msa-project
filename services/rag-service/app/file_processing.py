import io
from pathlib import Path

from fastapi import HTTPException
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx", ".pptx"}


def _extract_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=400,
            detail="TXT files must be UTF-8 encoded.",
        ) from exc


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n".join(page for page in pages if page)


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    lines = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n".join(lines)


def _extract_pptx(file_bytes: bytes) -> str:
    presentation = Presentation(io.BytesIO(file_bytes))
    lines: list[str] = []
    for slide_idx, slide in enumerate(presentation.slides, start=1):
        slide_lines: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                slide_lines.append(shape.text.strip())
        if slide_lines:
            lines.append(f"Slide {slide_idx}")
            lines.extend(slide_lines)
    return "\n".join(lines)


def extract_text_from_file(filename: str, file_bytes: bytes) -> tuple[str, str]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Unsupported file type '{ext}'. "
                f"Supported: {sorted(SUPPORTED_EXTENSIONS)}"
            ),
        )

    if ext == ".txt":
        text = _extract_txt(file_bytes)
    elif ext == ".pdf":
        text = _extract_pdf(file_bytes)
    elif ext == ".docx":
        text = _extract_docx(file_bytes)
    else:
        text = _extract_pptx(file_bytes)

    cleaned = text.strip()
    if not cleaned:
        raise HTTPException(
            status_code=400,
            detail=f"No readable text found in '{filename}'.",
        )
    return cleaned, ext
